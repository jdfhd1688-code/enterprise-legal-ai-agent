from __future__ import annotations

import unittest
from io import BytesIO
from pathlib import Path

from docx import Document

from app.deliverables.docx_redline import DocumentRedlineEngine
from app.persistence.deliverable_repository import DeliverableRepository
from app.schemas.document import DocumentAnchor
from app.schemas.risk import Finding, RedlineSuggestion, ReviewDecision, Severity
from app.services.analysis_service import AnalysisService, InvalidFileError
from app.skills.document_parser import DocumentParserSkill
from tests.helpers import TempSettings


ROOT = Path(__file__).resolve().parents[1]
FIXTURE = ROOT / "tests" / "fixtures" / "demo_contract.docx"


def docx_bytes(*paragraphs: str) -> bytes:
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def finding(change_type: str, original: str = "原条款", suggested: str = "建议条款", index: int = 0) -> Finding:
    anchor = DocumentAnchor(document_id="demo", paragraph_index=index, clause_id="P0000", text=original)
    redline = RedlineSuggestion(
        risk_id="P0000", anchor=anchor, original_clause=original, suggested_clause=suggested,
        change_reason="测试", change_type=change_type, confidence=.9,
        approved_by_human=True, human_action="accept", human_final_clause=suggested,
    )
    return Finding(clause_id="P0000", risk_type="test", severity="high", issue="测试", redline=redline)


class DocxAnchorTests(unittest.TestCase):
    def test_docx_parser_preserves_paragraph_heading_and_table_anchors(self) -> None:
        parsed = DocumentParserSkill().parse(FIXTURE.name, FIXTURE.read_bytes())
        self.assertGreaterEqual(len(parsed.anchors), 10)
        self.assertTrue(any(a.heading == "第二条 付款条件" and "90日" in a.text for a in parsed.anchors))
        self.assertTrue(any(a.anchor_type == "table_cell" for a in parsed.anchors))

    def test_docx_chunks_keep_anchor_mapping(self) -> None:
        from app.skills.chunker import ChunkerSkill
        parsed = DocumentParserSkill().parse(FIXTURE.name, FIXTURE.read_bytes())
        chunks = ChunkerSkill().chunk(parsed)
        self.assertTrue(all(chunk.anchor is not None for chunk in chunks))
        self.assertTrue(any(chunk.anchor.paragraph_index == 5 for chunk in chunks))


class RedlineEngineTests(unittest.TestCase):
    def test_replace_redline(self) -> None:
        clean, _, result = DocumentRedlineEngine().generate(docx_bytes("原条款", "保留条款"), [finding("replace")])
        output = Document(BytesIO(clean))
        self.assertEqual(output.paragraphs[0].text, "建议条款")
        self.assertEqual(output.paragraphs[1].text, "保留条款")
        self.assertEqual(result[0].status, "applied")

    def test_insert_redline(self) -> None:
        clean, _, result = DocumentRedlineEngine().generate(docx_bytes("原条款"), [finding("insert")])
        self.assertEqual([p.text for p in Document(BytesIO(clean)).paragraphs], ["原条款", "建议条款"])
        self.assertEqual(result[0].status, "applied")

    def test_delete_redline(self) -> None:
        clean, _, result = DocumentRedlineEngine().generate(docx_bytes("原条款", "保留条款"), [finding("delete")])
        self.assertEqual([p.text for p in Document(BytesIO(clean)).paragraphs], ["保留条款"])
        self.assertEqual(result[0].status, "applied")

    def test_unapproved_redline_is_blocked(self) -> None:
        item = finding("replace")
        item.redline.approved_by_human = False
        _, _, result = DocumentRedlineEngine().generate(docx_bytes("原条款"), [item])
        self.assertEqual(result[0].reason, "redline_not_approved")

    def test_ambiguous_anchor_is_not_forced(self) -> None:
        item = finding("replace", original="甲方应支付全部合同款项")
        item.redline.anchor.text = "甲方应当支付全部合同款项"
        _, _, result = DocumentRedlineEngine().generate(
            docx_bytes("甲方应当支付全部合同款项", "甲方应当支付全部合同款项"), [item]
        )
        self.assertEqual(result[0].reason, "anchor_ambiguous")

    def test_table_anchor_requires_manual_redline(self) -> None:
        item = finding("replace")
        item.redline.anchor.anchor_type = "table_cell"
        _, _, result = DocumentRedlineEngine().generate(docx_bytes("原条款"), [item])
        self.assertEqual(result[0].status, "manual_required")


class FinalizationAndPersistenceTests(unittest.TestCase):
    def create_task(self, service: AnalysisService):
        return service.create_task(FIXTURE.name, FIXTURE.read_bytes(), "审查高风险条款")

    def test_human_accept_edit_reject_are_persisted_and_audited(self) -> None:
        with TempSettings(use_real_kb=True) as settings:
            service = AnalysisService(settings)
            task = self.create_task(service)
            mods = {}
            for index, item in enumerate(task.risk.findings):
                action = "accept" if index == 0 else "edit" if index == 1 else "reject"
                mods[item.clause_id] = {"redline_action": action, "human_final_clause": "人工编辑后的条款"}
            reviewed = service.submit_review(task.task_id, ReviewDecision.approve, modifications=mods, reviewer="demo_reviewer")
            self.assertTrue(reviewed.review_finalized)
            self.assertIn("accept", {i.action.value for i in reviewed.review_items})
            self.assertIn("edit", {i.action.value for i in reviewed.review_items})
            self.assertIn("reject", {i.action.value for i in reviewed.review_items})
            self.assertGreaterEqual(sum(e.event_type == "human_review_action" for e in reviewed.audit_events), len(reviewed.risk.findings))

    def test_pending_high_risk_blocks_finalization(self) -> None:
        with TempSettings(use_real_kb=True) as settings:
            service = AnalysisService(settings)
            task = self.create_task(service)
            with self.assertRaises(InvalidFileError):
                service.finalize_review(task.task_id)

    def test_docx_and_report_deliverables_have_hashes(self) -> None:
        with TempSettings(use_real_kb=True) as settings:
            service = AnalysisService(settings)
            task = self.create_task(service)
            reviewed = service.submit_review(task.task_id, ReviewDecision.approve)
            self.assertEqual({d.type for d in reviewed.deliverables}, {"final_report", "reviewed_contract", "redline_contract"})
            self.assertTrue(all(len(d.sha256) == 64 and d.size_bytes > 0 for d in reviewed.deliverables))
            for kind in ("final_report", "reviewed_contract", "redline_contract"):
                _, content = service.get_deliverable(reviewed.task_id, kind)
                self.assertGreater(len(content), 1000)

    def test_task_restores_without_reanalysis(self) -> None:
        with TempSettings(use_real_kb=True) as settings:
            service = AnalysisService(settings)
            task = self.create_task(service)
            reviewed = service.submit_review(task.task_id, ReviewDecision.approve)
            restored = AnalysisService(settings).get_task(reviewed.task_id)
            self.assertEqual(restored.final_status, "completed")
            self.assertEqual(len(restored.deliverables), 3)
            self.assertIsNone(restored.parsed_document)
            self.assertEqual(restored.chunks, [])

    def test_safe_filename_and_task_id_guard(self) -> None:
        with TempSettings(use_real_kb=True) as settings:
            repository = DeliverableRepository(settings)
            with self.assertRaises(ValueError):
                repository.write("../bad", "report", "x.docx", b"x")
            with self.assertRaises(ValueError):
                repository.write("TASK-1", "report", "../x.docx", b"x")

    def test_other_contract_paragraphs_survive_writeback(self) -> None:
        with TempSettings(use_real_kb=True) as settings:
            service = AnalysisService(settings)
            task = self.create_task(service)
            reviewed = service.submit_review(task.task_id, ReviewDecision.approve)
            _, content = service.get_deliverable(reviewed.task_id, "reviewed_contract")
            text = "\n".join(p.text for p in Document(BytesIO(content)).paragraphs)
            self.assertIn("第一条 服务内容", text)
            self.assertIn("乙方按照采购清单", text)


if __name__ == "__main__":
    unittest.main()
