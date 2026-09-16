"""Generate synthetic DOCX fixtures for Phase 4 acceptance."""

from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]


def build(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading("企业服务采购合同", 0)
    document.add_paragraph("本合同为 DEMO SAMPLE，仅用于产品验收，不包含真实企业或个人信息。")
    document.add_heading("第一条 服务内容", level=1)
    document.add_paragraph("乙方按照采购清单向甲方提供企业软件实施与技术支持服务。")
    document.add_heading("第二条 付款条件", level=1)
    document.add_paragraph("甲方应在收到乙方合法有效发票后90日内支付全部合同款项。")
    document.add_heading("第三条 违约责任", level=1)
    document.add_paragraph("乙方应承担因履行本合同产生的全部损失，包括全部间接损失，且责任不设上限。")
    document.add_heading("第四条 知识产权", level=1)
    document.add_paragraph("项目成果及相关知识产权全部归乙方所有，甲方仅可在合同期内使用。")
    document.add_heading("第五条 争议解决", level=1)
    document.add_paragraph("双方发生争议时提交乙方所在地有管辖权的人民法院处理。")
    table = document.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "验收期限"
    table.cell(0, 1).text = "验收方式"
    table.cell(1, 0).text = "3日"
    table.cell(1, 1).text = "逾期视为自动验收"
    document.save(path)


if __name__ == "__main__":
    build(ROOT / "data" / "demo_contracts" / "demo_phase4_contract.docx")
    build(ROOT / "tests" / "fixtures" / "demo_contract.docx")
