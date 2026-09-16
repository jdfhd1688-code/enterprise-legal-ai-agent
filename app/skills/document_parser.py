"""Document parsing skill."""

from __future__ import annotations

from pathlib import Path

from app.schemas.document import DocumentAnchor, DocumentPage, ParsedDocument
from app.tools.pdf_tool import PdfTool


class DocumentParserError(RuntimeError):
    pass


class DocumentParserSkill:
    """Parses PDF, TXT and DOCX into page-level text records."""

    allowed_extensions = {".pdf", ".txt", ".docx"}

    def parse(self, filename: str, data: bytes) -> ParsedDocument:
        extension = Path(filename).suffix.lower()
        if extension not in self.allowed_extensions:
            raise DocumentParserError(
                f"不支持的文件类型 {extension or '(无扩展名)'}。MVP 支持 PDF，并预留 TXT/DOCX。"
            )
        if not data:
            raise DocumentParserError("上传文件为空。")

        content_type = extension.lstrip(".") or "unknown"
        warnings: list[str] = []
        anchors: list[DocumentAnchor] = []
        if extension == ".pdf":
            pages = self._parse_pdf(data, filename)
        elif extension == ".txt":
            pages = self._parse_txt(data, filename)
        else:
            pages, anchors = self._parse_docx(data, filename)

        real_pages = [page for page in pages if page.text.strip()]
        if not real_pages:
            raise DocumentParserError("文档中没有可提取的文本内容。")
        total_length = sum(len(page.text) for page in pages)
        if total_length < 12:
            warnings.append("文档文本过短，分析结果可能不完整。")
        if len(real_pages) < len(pages):
            warnings.append("部分页面没有可提取文本，已跳过空页。")

        return ParsedDocument(
            source=filename,
            title=Path(filename).stem,
            pages=real_pages,
            file_size=len(data),
            content_type=content_type,
            parse_warnings=warnings,
            num_pages=len(real_pages),
            text_length=total_length,
            anchors=anchors,
        )

    @staticmethod
    def _parse_pdf(data: bytes, filename: str) -> list[DocumentPage]:
        try:
            pdf_pages = PdfTool.extract(data)
        except Exception as exc:  # noqa: BLE001 - convert tool error to user message
            raise DocumentParserError(str(exc)) from exc
        return [
            DocumentPage(source=filename, page_no=page.page_no, text=page.text, section="page")
            for page in pdf_pages
        ]

    @staticmethod
    def _parse_txt(data: bytes, filename: str) -> list[DocumentPage]:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("gb18030", errors="replace")
        return [DocumentPage(source=filename, page_no=1, text=text.strip(), section="page")]

    @staticmethod
    def _parse_docx(data: bytes, filename: str) -> tuple[list[DocumentPage], list[DocumentAnchor]]:
        try:
            import docx  # type: ignore
        except ImportError as exc:
            raise DocumentParserError("DOCX 解析需要 python-docx，请先安装依赖。") from exc
        try:
            from io import BytesIO

            document = docx.Document(BytesIO(data))
            document_id = Path(filename).stem
            anchors: list[DocumentAnchor] = []
            heading = ""
            offset = 0
            for index, paragraph in enumerate(document.paragraphs):
                text = paragraph.text.strip()
                if not text:
                    continue
                style_name = paragraph.style.name if paragraph.style else ""
                if style_name.lower().startswith("heading") or style_name in {"标题", "Title"}:
                    heading = text
                anchors.append(DocumentAnchor(
                    document_id=document_id, paragraph_index=index, clause_id=f"P{index:04d}",
                    heading=heading, text=text, anchor_type="paragraph",
                    source_span=(offset, offset + len(text)),
                ))
                offset += len(text) + 1
            table_offset = len(document.paragraphs)
            for table_index, table in enumerate(document.tables):
                for row_index, row in enumerate(table.rows):
                    for cell_index, cell in enumerate(row.cells):
                        text = "\n".join(p.text.strip() for p in cell.paragraphs if p.text.strip())
                        if not text:
                            continue
                        anchors.append(DocumentAnchor(
                            document_id=document_id, paragraph_index=table_offset,
                            clause_id=f"T{table_index:03d}R{row_index:03d}C{cell_index:03d}",
                            heading=heading, text=text, anchor_type="table_cell",
                            table_index=table_index, row_index=row_index, cell_index=cell_index,
                            source_span=(offset, offset + len(text)),
                        ))
                        table_offset += 1
                        offset += len(text) + 1
            text = "\n\n".join(anchor.text for anchor in anchors)
            return [DocumentPage(source=filename, page_no=1, text=text, section="document")], anchors
        except Exception as exc:  # noqa: BLE001
            raise DocumentParserError(f"DOCX 解析失败：{exc}") from exc
