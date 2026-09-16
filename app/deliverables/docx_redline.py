"""Portfolio-safe clean and highlighted DOCX redline generation."""

from __future__ import annotations

import re
from difflib import SequenceMatcher
from io import BytesIO

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.shared import RGBColor

from app.schemas.deliverable import RedlineApplyResult
from app.schemas.risk import Finding, RedlineSuggestion


def normalize_text(value: str) -> str:
    return re.sub(r"[\s\u3000，。；：、,.!！?？（）()]+", "", value or "").lower()


class RedlineApplyGuard:
    def validate(self, document, finding: Finding, seen: set[str]) -> tuple[bool, str]:
        redline = finding.redline
        risk_id = finding.clause_id
        if redline is None or redline.anchor is None:
            return False, "anchor_missing"
        if risk_id in seen:
            return False, "duplicate_redline"
        if redline.anchor.anchor_type != "paragraph":
            return False, "table_cell_anchor_not_supported"
        if redline.human_action not in {"accept", "edit", "auto_approved"} or not redline.approved_by_human:
            return False, "redline_not_approved"
        if redline.confidence < 0.70:
            return False, "confidence_below_apply_threshold"
        index = redline.anchor.paragraph_index
        if index < 0 or index >= len(document.paragraphs):
            return False, "anchor_not_found"
        paragraph_text = document.paragraphs[index].text
        expected = normalize_text(redline.original_clause)
        actual = normalize_text(paragraph_text)
        duplicates = sum(normalize_text(p.text) == actual and bool(actual) for p in document.paragraphs)
        if duplicates > 1 and expected != actual:
            return False, "anchor_ambiguous"
        if expected != actual and SequenceMatcher(None, expected, actual).ratio() < 0.92:
            return False, "anchor_text_mismatch"
        return True, "verified"


class DocumentRedlineEngine:
    def __init__(self) -> None:
        self.guard = RedlineApplyGuard()

    def generate(self, original: bytes, findings: list[Finding]) -> tuple[bytes, bytes, list[RedlineApplyResult]]:
        clean = Document(BytesIO(original))
        marked = Document(BytesIO(original))
        seen: set[str] = set()
        results: list[RedlineApplyResult] = []
        applicable = sorted(
            [finding for finding in findings if finding.redline and finding.redline.anchor],
            key=lambda item: item.redline.anchor.paragraph_index,
            reverse=True,
        )
        for finding in applicable:
            ok, reason = self.guard.validate(clean, finding, seen)
            index = finding.redline.anchor.paragraph_index
            if not ok:
                finding.redline.redline_status = "manual_required" if "table_cell" in reason or "ambiguous" in reason else "apply_failed"
                finding.redline.failure_reason = reason
                results.append(RedlineApplyResult(risk_id=finding.clause_id, status=finding.redline.redline_status, reason=reason, paragraph_index=index))
                continue
            self._apply_clean(clean, index, finding.redline)
            self._apply_marked(marked, index, finding.redline)
            finding.redline.redline_status = "applied"
            seen.add(finding.clause_id)
            results.append(RedlineApplyResult(risk_id=finding.clause_id, status="applied", reason="guard_verified", paragraph_index=index))
        return self._bytes(clean), self._bytes(marked), results

    @staticmethod
    def _final_text(redline: RedlineSuggestion) -> str:
        return redline.human_final_clause or redline.suggested_clause

    def _apply_clean(self, document, index: int, redline: RedlineSuggestion) -> None:
        paragraph = document.paragraphs[index]
        if redline.change_type == "replace":
            paragraph.text = self._final_text(redline)
        elif redline.change_type == "insert":
            self._insert_after(paragraph, self._final_text(redline))
        elif redline.change_type == "delete":
            paragraph._element.getparent().remove(paragraph._element)

    def _apply_marked(self, document, index: int, redline: RedlineSuggestion) -> None:
        paragraph = document.paragraphs[index]
        if redline.change_type in {"replace", "delete"}:
            self._clear(paragraph)
            old = paragraph.add_run(redline.original_clause)
            old.font.strike = True
            old.font.color.rgb = RGBColor(160, 50, 45)
        if redline.change_type == "replace":
            new = paragraph.add_run("\n" + self._final_text(redline))
            new.font.color.rgb = RGBColor(26, 112, 73)
            new.font.highlight_color = WD_COLOR_INDEX.BRIGHT_GREEN
        elif redline.change_type == "insert":
            inserted = self._insert_after(paragraph, self._final_text(redline))
            run = inserted.runs[0]
            run.font.color.rgb = RGBColor(26, 112, 73)
            run.font.highlight_color = WD_COLOR_INDEX.BRIGHT_GREEN

    @staticmethod
    def _clear(paragraph) -> None:
        for run in list(paragraph.runs):
            paragraph._element.remove(run._element)

    @staticmethod
    def _insert_after(paragraph, text: str) -> Paragraph:
        node = OxmlElement("w:p")
        paragraph._p.addnext(node)
        inserted = Paragraph(node, paragraph._parent)
        inserted.style = paragraph.style
        inserted.add_run(text)
        return inserted

    @staticmethod
    def _bytes(document) -> bytes:
        stream = BytesIO()
        document.save(stream)
        return stream.getvalue()
