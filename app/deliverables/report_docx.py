"""Generate a professional downloadable final review report DOCX."""

from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from app.agent.prompts import DISCLAIMER
from app.schemas.task import TaskRecord


class ReportDocxGenerator:
    def generate(self, task: TaskRecord) -> bytes:
        if task.risk is None:
            raise ValueError("任务没有最终风险结论")
        document = Document()
        normal = document.styles["Normal"]
        normal.font.name = "Microsoft YaHei"
        normal.font.size = Pt(10.5)
        title = document.add_paragraph(style="Title")
        title.add_run("合同审查最终报告")
        document.add_paragraph(f"合同：{task.original_filename}")
        document.add_paragraph(f"任务编号：{task.task_id}")
        document.add_paragraph(f"最终状态：{task.final_status}　风险等级：{task.risk.risk_level.value.upper()}　置信度：{task.risk.confidence:.0%}")

        document.add_heading("执行摘要", level=1)
        document.add_paragraph(task.risk.summary)
        counts = {level: sum(f.severity.value == level for f in task.risk.findings) for level in ("high", "medium", "low")}
        summary = document.add_table(rows=2, cols=4)
        summary.style = "Table Grid"
        for cell, text in zip(summary.rows[0].cells, ["发现总数", "高风险", "中风险", "低风险"]):
            cell.text = text
            self._shade(cell, "183B70")
            for run in cell.paragraphs[0].runs:
                run.font.color.rgb = RGBColor(255, 255, 255)
        for cell, value in zip(summary.rows[1].cells, [len(task.risk.findings), counts["high"], counts["medium"], counts["low"]]):
            cell.text = str(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

        document.add_heading("风险明细与最终意见", level=1)
        decisions = {item.risk_id: item for item in task.review_items}
        for index, finding in enumerate(task.risk.findings, 1):
            document.add_heading(f"{index} {finding.issue}", level=2)
            decision = decisions.get(finding.clause_id)
            document.add_paragraph(f"风险等级：{finding.severity.value}　证据状态：{finding.evidence_status.value}　人工复核：{decision.action.value if decision else '无需人工复核'}")
            self._label(document, "合同证据", finding.contract_evidence)
            self._label(document, "法律依据", "；".join(f"{item.title} {item.article_no} ({item.citation_status.value})" for item in finding.legal_basis) or "证据不足")
            self._label(document, "企业 Playbook", "；".join(f"{item.title}｜标准 {item.expected}｜实际 {item.actual}｜{item.version}" for item in finding.playbook_evidence) or "无适用偏离")
            self._label(document, "判断理由", finding.reasoning or finding.issue)
            self._label(document, "处理建议", finding.recommendation)
            if finding.redline:
                final_clause = finding.redline.human_final_clause or finding.redline.suggested_clause
                self._label(document, "最终 Redline", final_clause if finding.redline.human_action not in {"reject", "resolved"} else "人工拒绝或已另行解决，未写回合同")

        document.add_heading("证据与引用状态", level=1)
        document.add_paragraph(f"证据状态：{task.risk.evidence_status.value}。引用校验：{task.risk.citation_validation_summary or '未提供'}。")
        document.add_heading("免责声明", level=1)
        document.add_paragraph(DISCLAIMER + " 内置法规与企业规则均为 DEMO/SAMPLE 数据，正式决策前应由专业人员核验。")
        stream = BytesIO()
        document.save(stream)
        return stream.getvalue()

    @staticmethod
    def _label(document, label: str, text: str) -> None:
        paragraph = document.add_paragraph()
        paragraph.add_run(label + "：").bold = True
        paragraph.add_run(text or "—")

    @staticmethod
    def _shade(cell, fill: str) -> None:
        properties = cell._tc.get_or_add_tcPr()
        shade = OxmlElement("w:shd")
        shade.set(qn("w:fill"), fill)
        properties.append(shade)
